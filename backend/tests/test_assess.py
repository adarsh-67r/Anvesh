import io

from app.assess.extract import extract_sections, kind_for
from app.assess.generate import pick_sections, validate_mcqs


def test_kind_detection():
    assert kind_for("notes.PDF", "application/octet-stream") == "pdf"
    assert kind_for("deck", "application/vnd.openxmlformats-officedocument.presentationml.presentation") == "pptx"
    assert kind_for("virus.exe", "application/octet-stream") is None


def test_pptx_sections_are_referenced_by_slide():
    from pptx import Presentation

    prs = Presentation()
    for body in ["Sampling frame lists every unit of the population from which the sample is drawn.", "short"]:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Sampling"
        s.placeholders[1].text = body
    buf = io.BytesIO(); prs.save(buf)
    sections = extract_sections("pptx", buf.getvalue())
    assert [s["ref"] for s in sections] == ["Slide 1"]  # slide 2 is too short to question
    assert "Sampling frame" in sections[0]["text"]


def test_docx_sections_follow_headings():
    from docx import Document

    d = Document()
    d.add_heading("Consumer Price Index", 1)
    d.add_paragraph("The CPI measures change over time in prices of a basket of goods and services consumed by households.")
    d.add_heading("Base year", 1)
    d.add_paragraph("The base year is the reference period against which index values are compared, currently 2012=100.")
    buf = io.BytesIO(); d.save(buf)
    refs = [s["ref"] for s in extract_sections("docx", buf.getvalue())]
    assert refs == ["Section: Consumer Price Index", "Section: Base year"]


def test_pick_sections_spreads_over_long_material():
    sections = [{"ref": f"Page {i}", "text": "x" * 3000} for i in range(1, 41)]
    chosen = pick_sections(sections, budget=30000)
    assert sum(len(s["text"]) for s in chosen) <= 30000 + 3000
    assert chosen[0]["ref"] == "Page 1" and int(chosen[-1]["ref"].split()[1]) > 30


def test_validate_drops_bad_questions():
    good = {"question": "What does CPI measure?", "options": ["Prices", "Output", "Jobs", "Trade"], "answer": "Prices",
            "explanation": "See page 2.", "source": "Page 2", "competency": "price_statistics", "difficulty": 1}
    raw = [good, dict(good, question="dup", answer="Nope"), dict(good, options=["a", "a", "b", "c"]), "junk",
           dict(good, question="Second?", source="Page 99", competency="made_up")]
    out = validate_mcqs(raw, {"Page 2"}, {"price_statistics"}, 10)
    assert [q["text"] for q in out] == ["What does CPI measure?", "Second?"]
    assert out[1]["source_ref"] == "" and out[1]["competency_id"] == "price_statistics"


def test_source_refs_match_loosely():
    q = {"question": "Q?", "options": ["a", "b", "c", "d"], "answer": "a"}
    refs = {"Page 12", "Slide 3", "Video 4:05"}
    for given, want in [("[Page 12]", "Page 12"), ("page 12", "Page 12"), ("Slide 3.", "Slide 3"), ("Video 4:05", "Video 4:05")]:
        assert validate_mcqs([dict(q, source=given)], refs, set(), 1)[0]["source_ref"] == want
